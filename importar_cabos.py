"""
Importa o catálogo de cabos da planilha 'LINHA DE TRANSMISSÃO 4.xlsm' para o banco.

Uso:   python importar_cabos.py "caminho/LINHA DE TRANSMISSÃO 4.xlsm"

Fontes na planilha
- aba 'Data Base'                  -> condutores, para-raios (aço) e OPGW
- aba 'Características dos Cabos'  -> cabo contrapeso (aço cobreado)

Regras
- Idempotente: reimportar atualiza pelo 'codigo', sem duplicar.
- Valores NUNCA são corrigidos. Inconsistências viram texto em `alerta_importacao`.
- `rmg_mm` (não existe na planilha) e `observacao` (livre) são preservados nas reimportações.
- Conversões: rtc_kn = kgf x 0,00980665; modulo_gpa = (kgf/mm²) x 0,00980665;
  dilatação '18,9 x 10^-6' -> 18.9 (10⁻⁶/°C).
"""
import re
import sys
import warnings

import db

USO_MAP = {"CONDUTOR": "CONDUTOR", "PARA-RAIOS": "PARA-RAIOS", "ÓPTICO": "OPGW", "OPTICO": "OPGW"}
ABA_BASE = "Data Base"
ABA_CARAC = "Características dos Cabos"
ARQ = "LINHA DE TRANSMISSÃO 4.xlsm"


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip()


def _num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def parse_dilatacao(v):
    """'18,9 x 10^-6' -> 18.9 (em 10⁻⁶/°C). Devolve None se o formato não for reconhecido."""
    m = re.match(r"^\s*([\d.,]+)\s*x\s*10\^-6\s*$", str(v or ""))
    return float(m.group(1).replace(",", ".")) if m else None


def _formacao(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return f"{int(v)} fios"
    return _norm(v)


def avaliar_alertas(c):
    """Verificações de consistência (não alteram valores). Retorna texto ou None."""
    a = []
    if c["r_cc_20"] is not None and c["r_ca_75"] is not None and c["r_ca_75"] < c["r_cc_20"]:
        a.append(f"R_CA 75 °C ({c['r_ca_75']:g}) menor que R_CC 20 °C ({c['r_cc_20']:g}) Ω/km: "
                 "incoerente; conferir com o catálogo do fabricante antes de usar.")
    if c["uso"] == "CONDUTOR" and c["r_ca_75"] is None:
        a.append("R_CA 75 °C ausente na planilha.")
    if c["uso"] == "CONDUTOR" and c["ampacidade_a"] is None:
        a.append("Ampacidade ausente na planilha.")
    if c["dilatacao_1e6"] is None and c["_dilatacao_txt"]:
        a.append(f"Coeficiente de dilatação em formato não reconhecido: {c['_dilatacao_txt']!r}.")
    return " ".join(a) or None


def _idx(headers, inicio=None, exato=None):
    for i, h in enumerate(headers):
        if (exato is not None and h == exato) or (inicio is not None and exato is None and h.startswith(inicio)):
            return i
    raise ValueError(f"Coluna não encontrada na aba '{ABA_BASE}': {exato or inicio!r}. A planilha mudou de layout?")


def ler_base(ws):
    linhas = list(ws.iter_rows(values_only=True))
    h = [_norm(x) for x in linhas[0]]
    ix = {
        "codigo": _idx(h, exato="Condutor Código Cruzado"), "uso": _idx(h, exato="Uso"),
        "tipo": _idx(h, exato="Tipo de cabo"), "nome": _idx(h, exato="Condutor"),
        "bitola": _idx(h, exato="Bitola"), "formacao": _idx(h, "Formação"),
        "secao_mm2": _idx(h, "Seção"), "diametro_mm": _idx(h, "Diâmetro Nominal"),
        "massa_kg_m": _idx(h, "Peso"), "rtc_kgf": _idx(h, "Carga de Ruptura"),
        "ampacidade_a": _idx(h, "Ampacidade"), "r_cc_20": _idx(h, "Resistência Elétrica CC"),
        "r_ca_75": _idx(h, "Resistência Elétrica CA"), "modulo_kgf_mm2": _idx(h, "Módulo de Elasticidade"),
        "dilat": _idx(h, "Coeficiente de Dilatação Linear"), "alfa_20c": _idx(h, "Coeficiente de dilatação linear por grau"),
        "curto_ka_s": _idx(h, "Capacidade de corrente de curto"), "fornecedor": _idx(h, "Fornecedor"),
        "diam_fio_al_mm": _idx(h, "Diâmetro do Fio de Alumínio"), "diam_fio_aco_mm": _idx(h, "Diâmetro do Fio de Aço"),
    }
    out = []
    for n, r in enumerate(linhas[1:], start=2):
        uso_raw = _norm(r[ix["uso"]]).upper()
        if not uso_raw:                      # linhas vazias (fórmulas sem resultado)
            continue
        if uso_raw not in USO_MAP:
            raise ValueError(f"Linha {n}: uso desconhecido {uso_raw!r}.")
        c = {
            "codigo": _norm(r[ix["codigo"]]), "uso": USO_MAP[uso_raw], "tipo": _norm(r[ix["tipo"]]) or None,
            "nome": _norm(r[ix["nome"]]) or None, "bitola": _norm(r[ix["bitola"]]) or None,
            "formacao": _formacao(r[ix["formacao"]]),
            "fornecedor": _norm(r[ix["fornecedor"]]) or None,
            "origem": f"planilha {ARQ} › {ABA_BASE} › linha {n}",
        }
        for k in ("secao_mm2", "diametro_mm", "massa_kg_m", "rtc_kgf", "ampacidade_a", "r_cc_20", "r_ca_75",
                  "modulo_kgf_mm2", "alfa_20c", "curto_ka_s", "diam_fio_al_mm", "diam_fio_aco_mm"):
            c[k] = _num(r[ix[k]])
        c["massa_kg_m"] = round(c["massa_kg_m"], 6) if c["massa_kg_m"] is not None else None   # tira ruído de ponto flutuante
        c["rtc_kn"] = c["rtc_kgf"] * db.KGF_PARA_KN if c["rtc_kgf"] is not None else None
        c["modulo_gpa"] = c["modulo_kgf_mm2"] * db.KGF_MM2_PARA_GPA if c["modulo_kgf_mm2"] is not None else None
        c["_dilatacao_txt"] = r[ix["dilat"]]
        c["dilatacao_1e6"] = parse_dilatacao(c["_dilatacao_txt"])
        out.append(c)
    return out


def ler_contrapeso(ws):
    """Bloco 'Características do Cabo Contrapeso' (rótulos na col. B, valores na col. C)."""
    vals, achou = {}, False
    for r in ws.iter_rows(values_only=True):
        b, cc = (r[1] if len(r) > 1 else None), (r[2] if len(r) > 2 else None)
        if _norm(b) == "Características do Cabo Contrapeso":
            achou = True
            continue
        if achou and b is not None:
            vals[_norm(b)] = cc
    if not achou:
        return None
    def v(prefixo):
        """Valor pelo início do rótulo (evita depender de símbolos como Ω, ° ou ²)."""
        achados = [x for k, x in vals.items() if k.startswith(prefixo)]
        if len(achados) != 1:
            raise ValueError(f"Rótulo '{prefixo}…' encontrado {len(achados)}x em '{ABA_CARAC}' (esperado 1).")
        return achados[0]

    ruptura_dan = _num(v("Carga de Ruptura"))
    peso_kg_km = _num(v("Peso líquido nominal"))
    nfios = _num(v("N"))
    c = {
        "codigo": _norm(v("Cabo")), "uso": "CONTRAPESO", "tipo": "AÇO COBREADO",
        "nome": _norm(v("Tipo / Condutividade")) or None, "bitola": None,
        "formacao": f"{int(nfios)} fio(s)" if nfios else None,
        "secao_mm2": _num(v("Seção Efetiva")),          # seção efetiva (nominal = 21,2 mm²)
        "diametro_mm": _num(v("Diâmetro Externo")),
        "massa_kg_m": peso_kg_km / 1000.0 if peso_kg_km is not None else None,
        "rtc_kn": ruptura_dan / 100.0 if ruptura_dan is not None else None,   # 1 daN = 0,01 kN
        "ampacidade_a": None, "r_cc_20": _num(v("Resistência Elétrica")), "r_ca_75": None,
        "modulo_kgf_mm2": None, "dilatacao_1e6": None, "alfa_20c": None, "curto_ka_s": None,
        "fornecedor": None, "diam_fio_al_mm": None, "diam_fio_aco_mm": None, "_dilatacao_txt": None,
        "origem": f"planilha {ARQ} › {ABA_CARAC} › Cabo Contrapeso",
    }
    if not c["codigo"]:
        return None
    sec_nom = _num(v("Seção Nominal"))
    c["observacao_auto"] = f"Seção efetiva {c['secao_mm2']:g} mm² (nominal {sec_nom:g} mm²)." if sec_nom else None
    return c


def importar(caminho):
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(caminho, data_only=True)
    cabos = ler_base(wb[ABA_BASE])
    cp = ler_contrapeso(wb[ABA_CARAC]) if ABA_CARAC in wb.sheetnames else None
    if cp:
        cabos.append(cp)

    codigos = [c["codigo"] for c in cabos]
    dup = {x for x in codigos if codigos.count(x) > 1}
    if dup:
        raise ValueError(f"Códigos duplicados na planilha: {sorted(dup)}")

    db.init_db()
    res = {"inseridos": 0, "atualizados": 0, "alertas": []}
    for c in cabos:
        c["alerta_importacao"] = avaliar_alertas(c)
        obs_auto = c.pop("observacao_auto", None)
        c.pop("_dilatacao_txt", None)
        if obs_auto:                                     # só no contrapeso; não sobrescreve observação do usuário
            c["observacao"] = obs_auto
        r = db.salvar_cabo(c, preservar=("rmg_mm", "observacao"))
        res["inseridos" if r == "inserido" else "atualizados"] += 1
        if c["alerta_importacao"]:
            res["alertas"].append((c["codigo"], c["alerta_importacao"]))
    res["total"] = len(cabos)
    return res


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    r = importar(sys.argv[1])
    print(f"Banco: {db.descricao_banco()}")
    print(f"Cabos lidos: {r['total']} · inseridos: {r['inseridos']} · atualizados: {r['atualizados']}")
    print(f"Alertas de qualidade: {len(r['alertas'])}")
    for cod, msg in r["alertas"]:
        print(f"  - {cod}: {msg}")
    print("RMG não consta da planilha: cadastre pelo app (Módulo 2 › Catálogo de cabos › Editar cabo).")
